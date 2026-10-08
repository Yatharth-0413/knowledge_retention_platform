"""Text extraction for uploaded documents.

Perfect fidelity is not the goal here (see the project README, section 4) —
just enough structured text for chunking and downstream topic/embedding work.
"""

import csv
import io

import openpyxl
import pymupdf
from docx import Document as DocxDocument

from app.documents.models import DocumentType


class ExtractionError(Exception):
    pass


def extract_text(file_bytes: bytes, file_type: DocumentType) -> str:
    try:
        if file_type == DocumentType.PDF:
            return _extract_pdf(file_bytes)
        if file_type == DocumentType.DOCX:
            return _extract_docx(file_bytes)
        if file_type == DocumentType.XLSX:
            return _extract_xlsx(file_bytes)
        if file_type == DocumentType.CSV:
            return _extract_csv(file_bytes)
    except Exception as exc:  # noqa: BLE001 - surfaced to the caller as a processing failure
        raise ExtractionError(str(exc)) from exc
    raise ExtractionError(f"Unsupported file type: {file_type}")


def _extract_pdf(file_bytes: bytes) -> str:
    with pymupdf.open(stream=file_bytes, filetype="pdf") as pdf:
        pages = [page.get_text() for page in pdf]
    return "\n\n".join(pages).strip()


def _extract_docx(file_bytes: bytes) -> str:
    doc = DocxDocument(io.BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _xlsx_rows(file_bytes: bytes) -> list[list[str]]:
    """Each sheet's data rows, prefixed by a single-cell `["# Sheet: <title>"]`
    marker row. Shared by _extract_xlsx (flattened, for chunking/citations) and
    extract_rows (structured, for person-attribution matching - which filters the
    marker rows back out, since they're formatting, not content; see
    knowledge/person_attribution.py).
    """
    workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    rows: list[list[str]] = []
    for sheet in workbook.worksheets:
        rows.append([f"# Sheet: {sheet.title}"])
        for row in sheet.iter_rows(values_only=True):
            cells = [str(v) for v in row if v is not None]
            if cells:
                rows.append(cells)
    return rows


def _extract_xlsx(file_bytes: bytes) -> str:
    return "\n".join(" | ".join(row) for row in _xlsx_rows(file_bytes)).strip()


def _csv_rows(file_bytes: bytes) -> list[list[str]]:
    text = file_bytes.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    return [row for row in reader if any(cell.strip() for cell in row)]


def _extract_csv(file_bytes: bytes) -> str:
    return "\n".join(" | ".join(row) for row in _csv_rows(file_bytes)).strip()


def extract_rows(file_bytes: bytes, file_type: DocumentType) -> list[list[str]]:
    """Structured row access for XLSX/CSV, used by person-attribution matching
    (knowledge/person_attribution.py) - separate from extract_text's flattened
    chunking/citation text. Synthetic `# Sheet:` marker rows are filtered out
    since they're not real document content.
    """
    try:
        if file_type == DocumentType.XLSX:
            return [row for row in _xlsx_rows(file_bytes) if not (len(row) == 1 and row[0].startswith("# Sheet:"))]
        if file_type == DocumentType.CSV:
            return _csv_rows(file_bytes)
    except Exception as exc:  # noqa: BLE001 - surfaced to the caller as a processing failure
        raise ExtractionError(str(exc)) from exc
    raise ExtractionError(f"Row extraction not supported for: {file_type}")
